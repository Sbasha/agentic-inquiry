"""Isolated parser contract checks, distinct from real-document acceptance."""

import io
import zipfile
from pathlib import Path

import pytest

from agentic_inquiry.ingestion import capabilities, extract


def test_actual_parser_module_handles_a_moderate_source_file():
    path = Path(__file__).parents[1] / "src" / "agentic_inquiry" / "ingestion.py"
    result = extract(path, path.read_bytes())
    assert result["status"] == "ok", result
    assert len(result["symbols"]) > 10
    assert any(symbol["name"] == "extract" for symbol in result["symbols"])


def test_markdown_retains_original_heading_spans_and_does_not_open_source(tmp_path):
    path = tmp_path / "document.md"
    path.write_text("original must not change")
    supplied = b"# First\nSome text\n\n## Second\nEvidence\n"
    result = extract(path, supplied)
    assert result["status"] == "ok"
    assert result["segments"][1]["line_start"] == 4
    assert result["segments"][1]["location"]["heading"] == "Second"
    assert path.read_text() == "original must not change"
    assert "".join(segment["text"] for segment in result["segments"]) == supplied.decode()


def test_csv_multiline_fields_have_original_lines_and_row_coordinates():
    result = extract("rows.csv", b'first,second\n"multi\nline",value\n')
    assert result["status"] == "ok"
    row = result["segments"][1]
    assert row["line_start"] == 2 and row["line_end"] == 3
    assert row["location"]["row_start"] == 2
    assert row["location"]["cells"] == ["multi\nline", "value"]


def test_html_retains_headings_tables_and_ignores_scripts():
    result = extract(
        "page.html",
        b'<h1 id="title">Heading</h1><script>bad()</script><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>',
    )
    assert result["status"] == "ok"
    assert [segment["text"] for segment in result["segments"]] == ["Heading", "A | B", "1 | 2"]
    assert result["segments"][0]["location"]["id"] == "title"
    assert all("line_start" not in segment for segment in result["segments"])


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("a.pdf", b"not a pdf"),
        ("a.docx", b"not a document"),
        ("a.json", b'{"broken":'),
        ("a.xml", b"<broken>"),
        ("a.toml", b"key = ["),
        ("a.yaml", b"bad: ["),
    ],
)
def test_malformed_documents_have_explicit_safe_diagnostics(name, data):
    result = extract(name, data)
    assert result["status"] == "malformed", result
    assert result["segments"] == []
    assert data.decode() not in result.get("diagnostic", "")


def test_xml_entities_cannot_read_local_files(tmp_path):
    secret = tmp_path / "private.txt"
    secret.write_text("never disclose")
    source = f'<!DOCTYPE doc [<!ENTITY x SYSTEM "{secret.as_uri()}">]><doc>&x;</doc>'.encode()
    result = extract("document.xml", source)
    assert result["status"] == "malformed"
    assert "never disclose" not in str(result)


def test_yaml_streams_and_custom_tags_remain_searchable_without_construction():
    data = b"resource: !Ref example\n---\nsetting: another\n"
    result = extract("configuration.yaml", data)
    assert result["status"] == "ok"
    assert result["segments"][0]["text"] == data.decode()


def test_archive_expansion_is_bounded_before_document_parser():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<document>" + "x" * 10000 + "</document>")
    result = extract("archive.docx", output.getvalue(), {"max_archive_bytes": 1000})
    assert result["status"] == "resource_limit"
    assert "max_archive_bytes" in result["diagnostic"]


def test_time_output_file_and_syntax_bounds_are_explicit():
    assert extract("large.txt", b"large", {"max_file_bytes": 2})["status"] == "oversize"
    assert (
        extract("a.txt", b"word " * 5000, {"max_output_bytes": 2048})["status"] == "resource_limit"
    )
    assert (
        extract("a.py", b"def run():\n    return 1\n", {"max_nodes": 2})["status"]
        == "resource_limit"
    )
    assert extract("a.txt", b"ordinary", {"timeout_seconds": 0.001})["status"] == "timeout"


def test_unsupported_and_invalid_code_remain_distinct():
    assert extract("binary.exe", b"abc")["status"] == "unsupported"
    broken = extract("a.py", b"def broken(\n")
    assert broken["status"] == "partial"
    assert broken["coverage"]["structural"] == "invalid"
    fallback = extract("a.rb", b"def hello\n  puts 'hi'\nend\n")
    assert fallback["status"] == "partial"
    assert fallback["coverage"]["structural"] == "unsupported"
    assert fallback["segments"]


def test_ocr_is_explicit_and_exposes_prerequisite_without_network():
    from PIL import Image

    data = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(data, format="PNG")
    assert extract("scan.png", data.getvalue())["status"] == "image_only"
    status = extract("scan.png", data.getvalue(), {"ocr": True, "max_pixels": 10})["status"]
    assert status in {"resource_limit", "unavailable"}
    assert capabilities()["ocr"]["enabled_by_default"] is False


def test_office_locations_and_tables_use_document_coordinates():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    document = Document()
    document.add_heading("Heading", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Left", "Right"
    data = io.BytesIO()
    document.save(data)
    word = extract("notes.docx", data.getvalue())
    assert word["status"] == "ok"
    assert word["segments"][0]["location"]["paragraph"] == 1
    assert word["segments"][1]["text"] == "Left | Right"
    assert word["segments"][1]["location"]["table"] == 1

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(2))
    shape.text = "Slide evidence"
    data = io.BytesIO()
    presentation.save(data)
    slides = extract("slides.pptx", data.getvalue())
    assert slides["segments"][0]["location"]["slide"] == 1
    assert slides["segments"][0]["location"]["shape_id"] == shape.shape_id

    workbook = Workbook()
    workbook.active.title = "Costs"
    workbook.active["B3"], workbook.active["C3"] = 42, "=B3*2"
    data = io.BytesIO()
    workbook.save(data)
    sheet = extract("figures.xlsx", data.getvalue())
    assert sheet["status"] == "ok"
    assert sheet["segments"][0]["location"]["cell_start"] == "B3"
    assert "C3==B3*2" in sheet["segments"][0]["text"]
    assert sheet["segments"][0]["location"]["sheet"] == "Costs"
    for result in (word, slides, sheet):
        assert all("line_start" not in segment for segment in result["segments"])
