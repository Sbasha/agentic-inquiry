"""Local, resource-bounded extraction with format-specific citation locations."""

from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from importlib.metadata import PackageNotFoundError, version
from itertools import pairwise
from pathlib import Path

from .structure import LANGUAGES, StructuralLimitError, analyze

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".rst",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".csv",
    ".tsv",
    ".toml",
    ".ini",
    ".sql",
    ".sh",
    ".bash",
    ".zsh",
    ".css",
    ".scss",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".scala",
    ".tf",
    ".log",
    ".conf",
    ".cfg",
    *LANGUAGES,
}
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | DOCUMENT_SUFFIXES | IMAGE_SUFFIXES
DEFAULTS = {
    "max_file_bytes": 32 * 1024 * 1024,
    "timeout_seconds": 30,
    "max_output_bytes": 16 * 1024 * 1024,
    "memory_mb": 768,
    "max_archive_bytes": 128 * 1024 * 1024,
    "max_archive_members": 10000,
    "max_pages": 2000,
    "max_nodes": 200000,
    "max_pixels": 20000000,
    "max_depth": 128,
    "ocr": False,
    "ocr_language": "eng",
    "language_overrides": {},
}


class ExtractionFailure(Exception):
    def __init__(self, status: str, diagnostic: str):
        super().__init__(diagnostic)
        self.status, self.diagnostic = status, diagnostic


def _settings(settings: dict | None) -> dict:
    supplied = settings or {}
    unknown = set(supplied) - set(DEFAULTS)
    if unknown:
        raise ValueError(f"Unknown parser setting: {min(unknown)}")
    configured = {**DEFAULTS, **supplied}
    for key, default in DEFAULTS.items():
        if isinstance(default, int) and not isinstance(default, bool):
            value = configured[key]
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or value <= 0
            ):
                raise ValueError(f"{key} must be positive and finite")
            if key != "timeout_seconds" and not isinstance(value, int):
                raise ValueError(f"{key} must be an integer")
    if configured["memory_mb"] < 32:
        raise ValueError("memory_mb must be at least 32")
    if not isinstance(configured["ocr"], bool):
        raise TypeError("ocr must be a boolean")
    if not isinstance(configured["ocr_language"], str) or not re.fullmatch(
        r"[A-Za-z0-9_+]{1,80}", configured["ocr_language"]
    ):
        raise ValueError("Invalid OCR language identifier")
    overrides = configured["language_overrides"]
    if not isinstance(overrides, dict) or any(
        suffix not in TEXT_SUFFIXES or language not in LANGUAGES.values()
        for suffix, language in overrides.items()
    ):
        raise ValueError(
            "Language overrides must map supported source suffixes to structural languages"
        )
    return configured


def _result(path: Path, status: str = "ok", diagnostic: str | None = None) -> dict:
    result = {
        "format": path.suffix.lower().lstrip(".") or "text",
        "language": LANGUAGES.get(path.suffix.lower()),
        "status": status,
        "segments": [],
        "symbols": [],
        "relations": [],
        "coverage": {
            "complete": status in {"ok", "empty"},
            "structural": "not_applicable",
            "limitations": [],
        },
        "parser": {"extraction_version": 1},
    }
    if diagnostic:
        result["diagnostic"] = diagnostic
        result["coverage"]["limitations"].append(diagnostic)
    return result


def _rss_mb(group: int, directory: Path) -> float | None:
    """Darwin does not enforce RLIMIT_AS; also account for OCR child processes."""
    import psutil

    try:
        total = psutil.Process(group).memory_info().rss
        child_file = directory / "ocr.pid"
        if child_file.exists():
            try:
                total += psutil.Process(int(child_file.read_text())).memory_info().rss
            except (psutil.NoSuchProcess, FileNotFoundError, ValueError):
                pass
        return total / (1024 * 1024)
    except psutil.NoSuchProcess:
        return 0.0
    except (psutil.Error, OSError):
        return None


def extract(path: str | Path, data: bytes, settings: dict | None = None) -> dict:
    """Parse supplied bytes in an isolated process; never open or modify the source path.

    Locations refer to source lines/bytes only for original UTF-8 source. Document
    locations identify pages, paragraphs, shapes or cells, with extracted segment
    offsets separate from original-file coordinates. No parser downloads assets.
    """
    path, configured = Path(path), _settings(settings)
    if len(data) > configured["max_file_bytes"]:
        return _result(path, "oversize", "Source exceeds max_file_bytes")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES and path.name.lower() not in {
        "dockerfile",
        "makefile",
        "license",
        "notice",
    }:
        return _result(path, "unsupported", "Unsupported source format")
    with tempfile.TemporaryDirectory(prefix="agentic-inquiry-extract-") as temporary:
        directory = Path(temporary)
        source, request, destination = (
            directory / "input",
            directory / "request.json",
            directory / "result.json",
        )
        source.write_bytes(data)
        request.write_text(
            json.dumps({"path": str(path), "settings": configured}), encoding="utf-8"
        )
        process = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "agentic_inquiry.ingestion",
                "--worker",
                str(request),
                str(source),
                str(destination),
            ],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
        deadline, next_memory_check = time.monotonic() + configured["timeout_seconds"], 0.0
        failure = None
        try:
            while process.poll() is None:
                now = time.monotonic()
                if now >= deadline:
                    failure = ("timeout", "Extraction exceeded timeout_seconds")
                    break
                if now >= next_memory_check:
                    resident = _rss_mb(process.pid, directory)
                    if resident is None and sys.platform == "darwin":
                        failure = ("unavailable", "Cannot enforce parser memory limit on this host")
                        break
                    if resident is not None and resident > configured["memory_mb"]:
                        failure = ("resource_limit", "Extraction exceeded memory_mb")
                        break
                    next_memory_check = now + 0.1
                time.sleep(min(0.02, max(0, deadline - now)))
        finally:
            if process.poll() is None:
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            process.wait()
        if failure:
            return _result(path, *failure)
        if process.returncode != 0:
            return _result(
                path,
                "resource_limit"
                if process.returncode in {-signal.SIGXCPU, -signal.SIGXFSZ}
                else "failed",
                f"Parser process terminated before publishing extraction (exit {process.returncode})",
            )
        try:
            with destination.open("rb") as stream:
                output = stream.read(configured["max_output_bytes"] + 1)
            if len(output) > configured["max_output_bytes"]:
                return _result(path, "resource_limit", "Extraction exceeded max_output_bytes")
            result = json.loads(output)
        except (OSError, ValueError):
            return _result(path, "failed", "Parser did not produce a valid result")
        result["limits"] = {
            key: value
            for key, value in configured.items()
            if key.startswith("max_") or key in {"timeout_seconds", "memory_mb"}
        }
        return result


def _segment(text: str, location: dict, **extra) -> dict:
    return {"text": text, "location": location, **extra}


def _table(rows: list[list]) -> str:
    return "\n".join(
        " | ".join(str(cell).replace("\n", " ") if cell is not None else "" for cell in row)
        for row in rows
    )


def _archive(data: bytes, settings: dict) -> None:
    if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        import olefile

        with olefile.OleFileIO(io.BytesIO(data)) as container:
            if container.exists("EncryptionInfo") and container.exists("EncryptedPackage"):
                raise ExtractionFailure(
                    "encrypted", "Encrypted Office document; no decryption attempted"
                )
        raise ExtractionFailure("unsupported", "Legacy Office container is not an OOXML document")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > settings["max_archive_members"]:
            raise ExtractionFailure("resource_limit", "Office archive exceeds max_archive_members")
        if sum(member.file_size for member in members) > settings["max_archive_bytes"]:
            raise ExtractionFailure("resource_limit", "Office archive exceeds max_archive_bytes")
        from defusedxml.ElementTree import fromstring

        for member in members:
            if member.flag_bits & 1:
                raise ExtractionFailure("encrypted", "Encrypted Office archive")
            if Path(member.filename).is_absolute() or ".." in Path(member.filename).parts:
                raise ExtractionFailure(
                    "malformed", "Office archive contains an invalid member path"
                )
            if member.filename.endswith(".xml"):
                root = fromstring(archive.read(member))
                stack = [(root, 1)]
                while stack:
                    element, depth = stack.pop()
                    if depth > settings["max_depth"]:
                        raise ExtractionFailure("resource_limit", "XML exceeds max_depth")
                    stack.extend((child, depth + 1) for child in element)


def _plain(path: Path, data: bytes, settings: dict) -> dict:
    text = data.decode("utf-8")
    if "\0" in text:
        raise ExtractionFailure("unsupported", "Binary data is not UTF-8 source text")
    result = _result(path)
    suffix = path.suffix.lower()
    if suffix in LANGUAGES or suffix in settings["language_overrides"]:
        result.update(
            analyze(
                path,
                data,
                max_nodes=settings["max_nodes"],
                language=settings["language_overrides"].get(suffix),
            )
        )
        result["coverage"]["complete"] = result["coverage"]["structural"] == "parsed"
        if not result["coverage"]["complete"]:
            result["status"] = "partial"
        return result
    if suffix == ".json":
        json.loads(text)
    elif suffix in {".yaml", ".yml"}:
        import yaml

        # Compose syntax without constructing custom application tags or executing constructors.
        for _ in yaml.compose_all(text, Loader=yaml.SafeLoader):
            pass
    elif suffix == ".toml":
        import tomllib

        tomllib.loads(text)
    elif suffix == ".xml":
        from defusedxml.ElementTree import fromstring

        root = fromstring(data)
        stack = [(root, 1)]
        while stack:
            element, depth = stack.pop()
            if depth > settings["max_depth"]:
                raise ExtractionFailure("resource_limit", "XML exceeds max_depth")
            stack.extend((child, depth + 1) for child in element)
    if suffix in {".csv", ".tsv"}:
        csv.field_size_limit(settings["max_output_bytes"])
        lines = text.splitlines(keepends=True)
        reader = csv.reader(
            io.StringIO(text, newline=""), delimiter="\t" if suffix == ".tsv" else ",", strict=True
        )
        start = 0
        for row_number, row in enumerate(reader, 1):
            end = reader.line_num
            result["segments"].append(
                _segment(
                    "".join(lines[start:end]),
                    {
                        "kind": "csv",
                        "row_start": row_number,
                        "row_end": row_number,
                        "column_start": 1,
                        "column_end": len(row),
                        "cells": row,
                    },
                    line_start=start + 1,
                    line_end=end,
                )
            )
            start = end
    else:
        lines = text.splitlines(keepends=True)
        starts = [0]
        if suffix in {".md", ".rst"}:
            starts += [
                index for index, line in enumerate(lines) if index and re.match(r"^#{1,6}\s+", line)
            ]
        starts.append(len(lines))
        for start, end in pairwise(starts):
            content = "".join(lines[start:end])
            if content.strip():
                location = {
                    "kind": "text",
                    "line_start": start + 1,
                    "line_end": end,
                    "byte_start": len("".join(lines[:start]).encode("utf-8")),
                    "byte_end": len("".join(lines[:end]).encode("utf-8")),
                }
                if suffix == ".md" and re.match(r"^#{1,6}\s+", lines[start]):
                    location["heading"] = lines[start].strip().lstrip("#").strip()
                result["segments"].append(
                    _segment(content, location, line_start=start + 1, line_end=end)
                )
    if suffix in {".rb", ".php", ".swift", ".kt", ".scala", ".sql", ".sh", ".bash", ".zsh"}:
        result["coverage"].update(structural="unsupported", complete=False)
        result["coverage"]["limitations"].append(
            "Original source text is searchable; structural analysis is not activated for this language"
        )
        result["status"] = "partial"
    return result


def _html(path: Path, data: bytes) -> dict:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data.decode("utf-8"), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    result = _result(path)
    result["parser"].update(name="beautifulsoup4", version=version("beautifulsoup4"))
    tags = soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "tr"])
    heading = None
    for number, tag in enumerate(tags, 1):
        if tag.name != "tr" and tag.find_parent(["p", "li", "pre", "tr"]):
            continue
        text = (
            _table(
                [
                    [
                        cell.get_text(" ", strip=True)
                        for cell in tag.find_all(["th", "td"], recursive=False)
                    ]
                ]
            )
            if tag.name == "tr"
            else tag.get_text(" ", strip=True)
        )
        if re.fullmatch(r"h[1-6]", tag.name):
            heading = text
        if text.strip():
            result["segments"].append(
                _segment(
                    text,
                    {
                        "kind": "html",
                        "element": tag.name,
                        "element_number": number,
                        "id": tag.get("id"),
                        "heading": heading,
                        "markup_line": tag.sourceline,
                    },
                )
            )
    if not result["segments"] and soup.get_text(" ", strip=True):
        result["segments"].append(
            _segment(soup.get_text(" ", strip=True), {"kind": "html", "element": "document"})
        )
    return result


def _docx(path: Path, data: bytes, settings: dict) -> dict:
    from docx import Document
    from docx.table import Table

    _archive(data, settings)
    result = _result(path)
    result["parser"].update(name="python-docx", version=version("python-docx"))
    document = Document(io.BytesIO(data))
    paragraph, table_number = 0, 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table_number += 1
            for row_number, row in enumerate(block.rows, 1):
                text = _table([[cell.text for cell in row.cells]])
                if text.strip():
                    result["segments"].append(
                        _segment(
                            text,
                            {
                                "kind": "docx",
                                "table": table_number,
                                "row_start": row_number,
                                "row_end": row_number,
                                "cell_start": 1,
                                "cell_end": len(row.cells),
                            },
                        )
                    )
        else:
            paragraph += 1
            if block.text.strip():
                result["segments"].append(
                    _segment(
                        block.text,
                        {
                            "kind": "docx",
                            "paragraph": paragraph,
                            "style": block.style.name if block.style else None,
                        },
                    )
                )
    # Headers and footers are separate source parts, never body paragraph numbers.
    seen = set()
    for section_number, section in enumerate(document.sections, 1):
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            part = getattr(section, name)
            if part.is_linked_to_previous:
                continue
            for number, paragraph in enumerate(part.paragraphs, 1):
                identity = (str(part.part.partname), number)
                if identity not in seen and paragraph.text.strip():
                    result["segments"].append(
                        _segment(
                            paragraph.text,
                            {
                                "kind": "docx",
                                "section": section_number,
                                "part": name,
                                "paragraph": number,
                            },
                        )
                    )
                    seen.add(identity)
    result["coverage"]["limitations"] = [
        "Text boxes, drawings, comments, footnotes and tracked revisions are not extracted"
    ]
    result["coverage"]["complete"] = False
    return result


def _pptx(path: Path, data: bytes, settings: dict) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    _archive(data, settings)
    result = _result(path)
    result["parser"].update(name="python-pptx", version=version("python-pptx"))
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > settings["max_pages"]:
        raise ExtractionFailure("resource_limit", "Presentation exceeds max_pages")
    for slide_number, slide in enumerate(presentation.slides, 1):
        stack = [(shape, 1) for shape in reversed(list(slide.shapes))]
        while stack:
            shape, depth = stack.pop()
            if depth > settings["max_depth"]:
                raise ExtractionFailure("resource_limit", "Grouped shapes exceed max_depth")
            location = {
                "kind": "pptx",
                "slide": slide_number,
                "shape_id": shape.shape_id,
                "shape_name": shape.name,
            }
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack.extend((child, depth + 1) for child in reversed(list(shape.shapes)))
            if shape.has_text_frame:
                for number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    if paragraph.text.strip():
                        result["segments"].append(
                            _segment(
                                paragraph.text,
                                {**location, "paragraph": number, "level": paragraph.level},
                            )
                        )
            if shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    result["segments"].append(
                        _segment(
                            _table([[cell.text for cell in row.cells]]),
                            {
                                **location,
                                "table": True,
                                "row_start": row_number,
                                "row_end": row_number,
                                "cell_start": 1,
                                "cell_end": len(row.cells),
                            },
                        )
                    )
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame and frame.text.strip():
                result["segments"].append(
                    _segment(
                        frame.text, {"kind": "pptx", "slide": slide_number, "part": "speaker_notes"}
                    )
                )
    result["coverage"]["limitations"] = [
        "Charts, drawings, embedded media and image text are not extracted"
    ]
    result["coverage"]["complete"] = False
    return result


def _xlsx(path: Path, data: bytes, settings: dict) -> dict:
    from openpyxl import load_workbook

    _archive(data, settings)
    result = _result(path)
    result["parser"].update(name="openpyxl", version=version("openpyxl"))
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > settings["max_pages"]:
            raise ExtractionFailure("resource_limit", "Workbook exceeds max_pages")
        cells_seen = 0
        for sheet in workbook.worksheets:
            # Discard untrusted dimension metadata and discover actual worksheet rows.
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                cells_seen += len(row)
                if cells_seen > settings["max_nodes"]:
                    raise ExtractionFailure("resource_limit", "Workbook exceeds max_nodes cells")
                populated = [cell for cell in row if cell.value is not None]
                if not populated:
                    continue
                text = _table([[f"{cell.coordinate}={cell.value}" for cell in populated]])
                result["segments"].append(
                    _segment(
                        text,
                        {
                            "kind": "xlsx",
                            "sheet": sheet.title,
                            "cell_start": populated[0].coordinate,
                            "cell_end": populated[-1].coordinate,
                            "row_start": populated[0].row,
                            "row_end": populated[-1].row,
                            "formula_policy": "source_formula_not_evaluated",
                        },
                    )
                )
    finally:
        workbook.close()
    result["coverage"]["limitations"] = [
        "Source formulas are retained without calculation; charts, drawings and cached formula results are not extracted"
    ]
    result["coverage"]["complete"] = False
    return result


def _ocr(image, location: dict, settings: dict, directory: Path) -> list[dict]:
    executable = shutil.which("tesseract")
    if not executable:
        raise ExtractionFailure(
            "unavailable", "Local OCR requires Tesseract and the selected language data"
        )
    if image.width * image.height > settings["max_pixels"]:
        raise ExtractionFailure("resource_limit", "OCR image exceeds max_pixels")
    source, output = directory / "ocr.png", directory / "ocr"
    image.save(source)
    command = [executable, str(source), str(output), "-l", settings["ocr_language"], "tsv"]
    completed = subprocess.Popen(
        command,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env={**os.environ, "OMP_THREAD_LIMIT": "1"},
    )
    child_file = directory / "ocr.pid"
    child_file.write_text(str(completed.pid))
    try:
        completed.wait(timeout=settings["timeout_seconds"])
    finally:
        if completed.poll() is None:
            completed.kill()
            completed.wait()
        child_file.unlink(missing_ok=True)
    if completed.returncode:
        raise ExtractionFailure(
            "unavailable", "Tesseract failed; check selected language data and local installation"
        )
    if output.with_suffix(".tsv").stat().st_size > settings["max_output_bytes"]:
        raise ExtractionFailure("resource_limit", "OCR exceeded max_output_bytes")
    groups = {}
    with output.with_suffix(".tsv").open(encoding="utf-8") as stream:
        for word in csv.DictReader(stream, delimiter="\t", quoting=csv.QUOTE_NONE):
            if not word.get("text", "").strip() or word.get("level") != "5":
                continue
            key = tuple(word[key] for key in ("page_num", "block_num", "par_num", "line_num"))
            groups.setdefault(key, []).append(word)
    segments = []
    for key, words in groups.items():
        left, top = min(int(w["left"]) for w in words), min(int(w["top"]) for w in words)
        right = max(int(w["left"]) + int(w["width"]) for w in words)
        bottom = max(int(w["top"]) + int(w["height"]) for w in words)
        segments.append(
            _segment(
                " ".join(w["text"] for w in words),
                {
                    **location,
                    "provenance": "ocr",
                    "ocr_engine": "tesseract",
                    "ocr_language": settings["ocr_language"],
                    "ocr_line": list(key),
                    "bbox": [left, top, right, bottom],
                    "coordinate_system": "rendered_image_pixels_top_left",
                    "render_width": image.width,
                    "render_height": image.height,
                    "recognition_confidence": sum(float(w["conf"]) for w in words) / len(words),
                },
            )
        )
    return segments


def _pdf(path: Path, data: bytes, settings: dict, directory: Path) -> dict:
    import pdfplumber

    result = _result(path)
    result["parser"].update(name="pdfplumber", version=version("pdfplumber"))
    missing_pages = []
    with pdfplumber.open(io.BytesIO(data)) as document:
        if len(document.pages) > settings["max_pages"]:
            raise ExtractionFailure("resource_limit", "PDF exceeds max_pages")
        for number, page in enumerate(document.pages, 1):
            tables = [
                table
                for table in page.find_tables()
                if len(table.rows) > 1 and len(table.columns) > 1
            ]
            boxes = [table.bbox for table in tables]
            for table_number, table in enumerate(tables, 1):
                rows = table.extract()
                for row_number, row in enumerate(rows, 1):
                    if any(cell for cell in row):
                        box = table.rows[row_number - 1].bbox
                        result["segments"].append(
                            _segment(
                                _table([row]),
                                {
                                    "kind": "pdf",
                                    "page": number,
                                    "table": table_number,
                                    "row_start": row_number,
                                    "row_end": row_number,
                                    "bbox": list(box),
                                    "coordinate_system": "points_top_left",
                                    "provenance": "embedded_text",
                                },
                            )
                        )
            for line_number, line in enumerate(page.extract_text_lines(), 1):
                center_x, center_y = (
                    (line["x0"] + line["x1"]) / 2,
                    (line["top"] + line["bottom"]) / 2,
                )
                if any(
                    x0 <= center_x <= x1 and top <= center_y <= bottom
                    for x0, top, x1, bottom in boxes
                ):
                    continue
                if line["text"].strip():
                    result["segments"].append(
                        _segment(
                            line["text"],
                            {
                                "kind": "pdf",
                                "page": number,
                                "region": line_number,
                                "bbox": [line["x0"], line["top"], line["x1"], line["bottom"]],
                                "coordinate_system": "points_top_left",
                                "provenance": "embedded_text",
                            },
                        )
                    )
            if not page.chars:
                if settings["ocr"]:
                    import pypdfium2

                    with pypdfium2.PdfDocument(data) as rendered:
                        pdf_page = rendered[number - 1]
                        scale = 2.0
                        if (
                            pdf_page.get_width() * pdf_page.get_height() * scale * scale
                            > settings["max_pixels"]
                        ):
                            raise ExtractionFailure(
                                "resource_limit", "Rendered PDF page exceeds max_pixels"
                            )
                        bitmap = pdf_page.render(scale=scale)
                        try:
                            recognized = _ocr(
                                bitmap.to_pil(),
                                {"kind": "pdf", "page": number, "render_scale": scale},
                                settings,
                                directory,
                            )
                            result["segments"].extend(recognized)
                            if not recognized:
                                missing_pages.append(number)
                        finally:
                            bitmap.close()
                            pdf_page.close()
                else:
                    missing_pages.append(number)
            page.close()
    result["coverage"].update(
        pages=number if "number" in locals() else 0, image_only_pages=missing_pages
    )
    if missing_pages:
        result["coverage"]["complete"] = False
        result["coverage"]["limitations"].append(
            "Pages without extractable text require local OCR; blank pages may also appear here"
        )
        result["status"] = "partial" if result["segments"] else "image_only"
    return result


def _image(path: Path, data: bytes, settings: dict, directory: Path) -> dict:
    from PIL import Image, ImageSequence

    result = _result(path)
    result["parser"].update(name="Pillow/Tesseract", pillow_version=version("Pillow"))
    if not settings["ocr"]:
        return _result(path, "image_only", "Enable local OCR to recognize image text")
    with Image.open(io.BytesIO(data)) as image:
        for number, frame in enumerate(ImageSequence.Iterator(image), 1):
            if number > settings["max_pages"]:
                raise ExtractionFailure("resource_limit", "Image exceeds max_pages frames")
            result["segments"].extend(
                _ocr(frame, {"kind": "image", "frame": number}, settings, directory)
            )
    if not result["segments"]:
        result["status"] = "image_only"
        result["coverage"]["complete"] = False
    return result


def _direct(path: Path, data: bytes, settings: dict, directory: Path) -> dict:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES or path.name.lower() in {
            "dockerfile",
            "makefile",
            "license",
            "notice",
        }:
            result = _plain(path, data, settings)
        elif suffix in {".html", ".htm"}:
            result = _html(path, data)
        elif suffix in IMAGE_SUFFIXES:
            result = _image(path, data, settings, directory)
        elif suffix == ".pdf":
            result = _pdf(path, data, settings, directory)
        else:
            result = {".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}[suffix](path, data, settings)
        if result["status"] == "ok" and not result["segments"]:
            result["status"] = "empty"
        return result
    except ExtractionFailure as error:
        return _result(path, error.status, error.diagnostic)
    except StructuralLimitError:
        return _result(path, "resource_limit", "Source exceeds max_nodes syntax nodes")
    except MemoryError:
        return _result(path, "resource_limit", "Parser exhausted its memory allowance")
    except ImportError as error:
        return _result(path, "unavailable", f"Parser dependency unavailable: {error.name}")
    except Exception as error:  # noqa: BLE001 - third-party parser errors stay inside the worker
        if type(error).__name__ in {"PDFPasswordIncorrect", "PDFEncryptionError"}:
            return _result(path, "encrypted", "Encrypted PDF; no decryption attempted")
        if isinstance(error, subprocess.TimeoutExpired):
            return _result(path, "timeout", "OCR exceeded timeout_seconds")
        # Parser exception messages can contain source text, names and payloads.
        return _result(path, "malformed", f"Parser rejected the source ({type(error).__name__})")


def capabilities() -> dict:
    """Describe actual supported coverage without loading embeddings or downloading."""
    dependencies = {}
    for package in (
        "pdfplumber",
        "python-docx",
        "python-pptx",
        "openpyxl",
        "beautifulsoup4",
        "tree-sitter-language-pack",
    ):
        try:
            dependencies[package] = version(package)
        except PackageNotFoundError:
            dependencies[package] = None
    return {
        "formats": sorted(SUPPORTED_SUFFIXES),
        "structural_languages": sorted(set(LANGUAGES.values())),
        "additional_bundled_grammars": "Not activated; other code remains original text with reduced structural coverage",
        "dependencies": dependencies,
        "ocr": {
            "engine": "tesseract",
            "available": shutil.which("tesseract") is not None,
            "enabled_by_default": False,
            "language_data": "Availability verified when OCR is requested",
        },
        "default_limits": DEFAULTS.copy(),
        "platforms": ["macOS", "Linux"],
        "relationship_resolution": "Syntax facts with unresolved runtime bindings",
    }


def _worker(request: Path, source: Path, destination: Path) -> None:
    import resource

    payload = json.loads(request.read_text(encoding="utf-8"))
    settings = _settings(payload["settings"])
    resource.setrlimit(
        resource.RLIMIT_CPU,
        (
            max(1, math.ceil(settings["timeout_seconds"])),
            max(2, math.ceil(settings["timeout_seconds"]) + 1),
        ),
    )
    resource.setrlimit(
        resource.RLIMIT_FSIZE, (settings["max_output_bytes"], settings["max_output_bytes"])
    )
    if sys.platform.startswith("linux"):
        memory = settings["memory_mb"] * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (memory, memory))
    result = _direct(Path(payload["path"]), source.read_bytes(), settings, source.parent)
    encoded = json.dumps(result, ensure_ascii=False).encode("utf-8")
    if len(encoded) > settings["max_output_bytes"]:
        encoded = json.dumps(
            _result(Path(payload["path"]), "resource_limit", "Extraction exceeded max_output_bytes")
        ).encode("utf-8")
    destination.write_bytes(encoded)


if __name__ == "__main__":
    if len(sys.argv) != 5 or sys.argv[1] != "--worker":
        raise SystemExit("Extraction worker is an internal operation")
    _worker(*(Path(argument) for argument in sys.argv[2:]))
